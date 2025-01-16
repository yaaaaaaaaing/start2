# -*- coding: utf-8 -*-
import requests
import json
import hashlib
import openpyxl
import crcmod
from send_receive_email import send_email
from init import *

def name_list_check(name_list_path):

    # 检查每个server email是否包含五个子服务
    # 检查是否有重复的client send_email
    # 返回最大行数
    name_list_wb = openpyxl.load_workbook(name_list_path)
    name_list_ws = name_list_wb["list"]
    client_vpn_list = []
    client_nf_list = []
    serv_email_conter = num_max_slots - 1
    for row in range(2,name_list_ws.max_row+1):
        serv_email = name_list_ws.cell(row=row,column=column_serv_email).value
        client_vpn_email = name_list_ws.cell(row=row,column=column_client_vpn_email).value
        client_nf_email = name_list_ws.cell(row=row,column=column_client_nf_email).value
        if serv_email:
            if serv_email_conter != num_max_slots - 1:
                print(f"serv_email {serv_email} is not in the right row")
                exit(1)
            elif serv_email == "#endofdata":
                serv_email_conter = 0
                max_row_list = row
                break
            else:
                serv_email_conter = 0
        else:
            serv_email_conter += 1
        
        

        if client_vpn_email:
            if client_vpn_email in client_vpn_list:
                print(f"client_vpn_email {client_vpn_email} is duplicated")
                exit(1)
            else:
                client_vpn_list.append(client_vpn_email)
        if client_nf_email:
            if client_nf_email in client_nf_list:
                print(f"client_nf_email {client_nf_email} is duplicated")
                exit(1)
            else:
                client_nf_list.append(client_nf_email)


    return max_row_list

def name_list_dict_gen(name_list_path,max_row_list):
    # 生成各类密码/pin
    name_list_wb = openpyxl.load_workbook(name_list_path)
    name_list_ws = name_list_wb["list"]
    for row in range(2,max_row_list):
        serv_email = name_list_ws.cell(row=row,column=column_serv_email).value
        client_vpn_email = name_list_ws.cell(row=row,column=column_client_vpn_email).value
        client_nf_email = name_list_ws.cell(row=row,column=column_client_nf_email).value

        crc16_func_email = crcmod.mkCrcFun(0x11021, initCrc=0, xorOut=0xFFFF, rev=True)
        crc16_func_nf = crcmod.mkCrcFun(0x18005, initCrc=0, xorOut=0xFFFF, rev=True)
        crc16_func_vpn = crcmod.mkCrcFun(0x13d65, initCrc=0, xorOut=0xFFFF, rev=True)
        crc16_func_pin = crcmod.mkCrcFun(0x18005, initCrc=0, xorOut=0xFFFF, rev=True)
        if serv_email is not None:
            serv_email_data = serv_email.encode('utf-8')
            name_list_ws.cell(row=row,column=column_serv_nf_pw).value = "PassWord~" + hex(crc16_func_nf(serv_email_data)).replace("0x","")
            counter = 0
        client_nf_email_data = serv_email_data+hex(counter).encode('utf-8')
        name_list_ws.cell(row=row,column=column_client_nf_pin).value = str(crc16_func_pin(client_nf_email_data)).zfill(4)[:4] + f" 位置{counter+1}"
        counter += 1
        
    name_list_wb.save(name_list_path)

def name_list_parse(name_list_path,max_row_list):
    name_list_wb = openpyxl.load_workbook(name_list_path)
    name_list_ws = name_list_wb["list"]
    email_info_list = []
    for row in range(2,max_row_list):
        serv_email = name_list_ws.cell(row=row,column=column_serv_email).value
        if serv_email is not None:
            email_info_dict = {}
            email_info_dict["serv_email"] = serv_email
            email_info_dict["serv_email_pw"] = name_list_ws.cell(row=row,column=column_serv_email_pw).value
            email_info_dict["serv_nf_pw"] = name_list_ws.cell(row=row,column=column_serv_nf_pw).value
            email_info_dict["serv_vpn_url"] = name_list_ws.cell(row=row,column=column_serv_vpn_pw).value
            email_info_dict["client_vpn_email"] = []
            email_info_dict["client_pin_email"] = []
            email_info_dict["client_nf_email"] = []
            client_vpn_email = name_list_ws.cell(row=row,column=column_client_vpn_email).value
            client_nf_pin = name_list_ws.cell(row=row,column=column_client_nf_pin).value
            client_nf_email = name_list_ws.cell(row=row,column=column_client_nf_email).value
            if client_vpn_email is not None:
                email_info_dict["client_vpn_email"].append(client_vpn_email)
            if client_nf_email is not None:
                email_info_dict["client_nf_email"].append(client_nf_email)
                email_info_dict["client_pin_email"].append(client_nf_pin)
            email_info_list.append(email_info_dict)
        else:
            email_info_dict = email_info_list[-1]
            client_vpn_email = name_list_ws.cell(row=row,column=column_client_vpn_email).value
            client_nf_pin = name_list_ws.cell(row=row,column=column_client_nf_pin).value
            client_nf_email = name_list_ws.cell(row=row,column=column_client_nf_email).value
            if client_vpn_email is not None:
                email_info_dict["client_vpn_email"].append(client_vpn_email)
            if client_nf_email is not None:
                email_info_dict["client_nf_email"].append(client_nf_email)
                email_info_dict["client_pin_email"].append(client_nf_pin)

    client_info_dict = {}
    for email_info_dict in email_info_list:
        for client_vpn_email in email_info_dict["client_vpn_email"]:
            if client_vpn_email not in client_info_dict:
                client_info_dict[client_vpn_email] = {}
                client_info_dict[client_vpn_email]["serv_vpn_email_pw"] = email_info_dict["serv_email_pw"]
                client_info_dict[client_vpn_email]["serv_vpn_email"] = email_info_dict["serv_email"]
                client_info_dict[client_vpn_email]["serv_vpn_url"] = email_info_dict["serv_vpn_url"]
            else:
                client_info_dict[client_vpn_email]["serv_vpn_email_pw"] = email_info_dict["serv_email_pw"]
                client_info_dict[client_vpn_email]["serv_vpn_email"] = email_info_dict["serv_email"]
                client_info_dict[client_vpn_email]["serv_vpn_url"] = email_info_dict["serv_vpn_url"]
        for client_nf_email in email_info_dict["client_nf_email"]:
            if client_nf_email not in client_info_dict:
                client_info_dict[client_nf_email] = {}
                client_info_dict[client_nf_email]["serv_nf_email_pw"] = email_info_dict["serv_email_pw"]
                client_info_dict[client_nf_email]["serv_nf_email"] = email_info_dict["serv_email"]
                client_info_dict[client_nf_email]["serv_nf_pw"] = email_info_dict["serv_nf_pw"]
                client_info_dict[client_nf_email]["serv_nf_pin"] = email_info_dict["client_pin_email"][email_info_dict["client_nf_email"].index(client_nf_email)]
            else:
                client_info_dict[client_nf_email]["serv_nf_email_pw"] = email_info_dict["serv_email_pw"]
                client_info_dict[client_nf_email]["serv_nf_email"] = email_info_dict["serv_email"]
                client_info_dict[client_nf_email]["serv_nf_pw"] = email_info_dict["serv_nf_pw"]
                client_info_dict[client_nf_email]["serv_nf_pin"] = email_info_dict["client_pin_email"][email_info_dict["client_nf_email"].index(client_nf_email)]

    return client_info_dict

def client_info_compare(client_info_dict,client_conf_path):
    with open(client_conf_path,'r') as client_conf_file:
        client_info_dict_old = json.load(client_conf_file)
    
    client_info_delta_dict = {}
    for key,item in client_info_dict.items():
        if key in client_info_dict_old:
            if item != client_info_dict_old[key]:
                client_info_delta_dict[key] = item
        else:
            client_info_delta_dict[key] = item
    

    with open(client_conf_path,'w') as client_conf_file:
        client_conf_file.write(json.dumps(client_info_dict, indent=2, ensure_ascii=False))
        client_conf_file.close()

    return client_info_delta_dict

def send_client_info(client_info_dict,release_email_dict):
    for client_email,infos in client_info_dict.items():
        body = ""
        attachment_files = []
        subject = "代理信息"
        release_info_list = release_email_dict["output client infos"]
        for key,item in infos.items():
            if key in release_info_list:
                body += f"{key}:\n{item}\n"
                body += "\n"
        for key,item in release_email_dict["attchment_mapping"].items():
            if key in infos:
                attachment_files.append(release_email_dict["attchment_mapping"][key])   
        send_email(subject,body,client_email,attachment_files)


def updata_name_list(name_list_path,client_conf_path,hash_path):
    max_row_list = name_list_check(name_list_path)
    name_list_dict_gen(name_list_path,max_row_list)
    client_info_dict = name_list_parse(name_list_path,max_row_list)
    client_info_delta_dict = client_info_compare(client_info_dict,client_conf_path)
    if client_info_delta_dict == {}:
        return False
    print("name list is updated")
    return True
    
def send_specific_info(input_email,client_conf_path):
    with open(client_conf_path,'r') as client_conf_file:
        client_info_dict = json.load(client_conf_file)
    if input_email in client_info_dict:
        output_dict = {}
        output_dict[input_email] = client_info_dict[input_email]
        send_client_info(output_dict,release_email_dict)



    

if __name__ == '__main__':

    print(1)
    


    